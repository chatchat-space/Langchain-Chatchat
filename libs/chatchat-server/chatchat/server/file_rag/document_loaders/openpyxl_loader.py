"""openpyxl 实现的 Excel 结构化加载器。

替代 UnstructuredExcelLoader（Windows 下段错误）。

设计：
- 表头识别：跳过开头的全空行，第一个非空行作为表头（要求 >=2 个非空单元格，
  否则视为无表头，退化为整表单文档模式）。
- 行转文档块：每个数据行生成一个独立 Document，
  page_content 为 "表头: 值" 的逐行拼接（自带上下文，利于向量化与检索），
  metadata 记录 source / sheet_name / row_id / title。
"""
from typing import Any, List, Optional

from langchain_core.documents import Document
from openpyxl import load_workbook


def _fmt_cell(v: Any) -> str:
    """单元格值转干净字符串：None→空串，浮点整数值去掉 .0，其余 str + strip。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (int,)):
        return str(v)
    return str(v).strip()


class OpenpyxlExcelLoader:
    def __init__(self, file_path: str, **kwargs):
        self.file_path = file_path

    def _sheet_to_docs(self, ws) -> List[Document]:
        # 1) 读取所有非全空行（记录 Excel 行号，从 1 开始）
        data_rows: List[tuple] = []  # (row_no, [cell_str, ...])
        for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_fmt_cell(c) for c in row]
            # 去掉行尾全空的多余列
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                data_rows.append((row_no, cells))

        if not data_rows:
            return []

        # 2) 表头识别：第一个非空行；非空单元格数 >=2 才认为是表头
        header_row_no, first = data_rows[0]
        header = [c for c in first]
        non_empty_header = [c for c in header if c != ""]
        has_header = len(non_empty_header) >= 2 and len(data_rows) > 1

        if not has_header:
            # 退化模式：无表头（或只有一行），整表拼一个文档
            text = "\n".join("\t".join(cells) for _, cells in data_rows)
            return [
                Document(
                    page_content=text,
                    metadata={
                        "source": self.file_path,
                        "sheet_name": ws.title,
                        "row_id": header_row_no,
                        "title": f"{ws.title}",
                    },
                )
            ]

        # 3) 行转文档块：每行一个 Document，"表头: 值" 拼接
        #    无名表头补 "列N"
        col_names: List[str] = []
        seen = 0
        for i, h in enumerate(header):
            col_names.append(h if h else f"列{i + 1}")
            seen += 1

        docs: List[Document] = []
        for row_no, cells in data_rows[1:]:
            # 行长度对齐表头
            cells = cells + [""] * (len(col_names) - len(cells))
            pairs = []
            for name, val in zip(col_names, cells):
                if val != "":
                    pairs.append(f"{name}: {val}")
            if not pairs:
                continue
            title = pairs[0][:50] if pairs else ""
            docs.append(
                Document(
                    page_content="\n".join(pairs),
                    metadata={
                        "source": self.file_path,
                        "sheet_name": ws.title,
                        "row_id": row_no,
                        "title": title,
                    },
                )
            )
        return docs

    def load(self) -> List[Document]:
        wb = load_workbook(self.file_path, read_only=True, data_only=True)
        docs: List[Document] = []
        try:
            for ws in wb.worksheets:
                docs.extend(self._sheet_to_docs(ws))
        finally:
            wb.close()
        return docs
